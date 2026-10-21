"""用 L6 示例实例验证通用化本体增强 / 实例导入 / 审批落库流程。

直接使用 FastAPI TestClient，不依赖外部 LLM/Neo4j/Jena，验证默认 JSONL 降级路径。
"""
from __future__ import annotations

import io
import json
import sys
import tempfile
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))
from main import app
from api import router_ontologies


def _mock_call_llm(*args, **kwargs):
    return {
        "choices": [{
            "message": {
                "content": json.dumps({
                    "name": "L6 采购计划通用本体（增强版）",
                    "description": "增加审批记录对象",
                    "objects": [
                        {
                            "id": "obj_approval_record",
                            "name": "审批记录",
                            "business_definition": "记录业务审批过程",
                            "attributes": [
                                {"id": "attr_approver", "name": "审批人", "data_type": "string", "required": False, "constraints": {}},
                                {"id": "attr_approval_time", "name": "审批时间", "data_type": "date", "required": False, "constraints": {}},
                                {"id": "attr_approval_comment", "name": "审批意见", "data_type": "string", "required": False, "constraints": {}},
                            ],
                            "relations": [],
                        }
                    ],
                    "constraints": [],
                }, ensure_ascii=False)
            }
        }]
    }


def _make_test_excel(tmpdir: Path) -> Path:
    """复制根目录 L6 示例实例到临时目录，避免污染工作区。"""
    src = Path(__file__).parent.parent / "L6采购计划_示例实例.xlsx"
    dst = tmpdir / "l6_instances.xlsx"
    dst.write_bytes(src.read_bytes())
    return dst


def _make_ontology_payload() -> dict:
    return {
        "name": "L6 采购计划通用本体",
        "description": "用于端到端验证的简化本体",
        "objects": [
            {
                "id": "obj_activity",
                "name": "活动",
                "business_definition": "采购计划流程中的业务活动",
                "stable_key": "attr_activity_code",
                "attributes": [
                    {"id": "attr_activity_code", "name": "活动编码", "data_type": "string", "required": True, "constraints": {}},
                    {"id": "attr_activity_name", "name": "活动名称", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_stage_name", "name": "阶段名称", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_leading_role", "name": "主导角色", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_participant_role", "name": "参与角色", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_input_doc", "name": "输入单据", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_output_doc", "name": "输出单据", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_description", "name": "描述", "data_type": "string", "required": False, "constraints": {}},
                ],
                "relations": [],
            },
            {
                "id": "obj_role",
                "name": "角色",
                "business_definition": "流程中的角色定义",
                "stable_key": "attr_role_code",
                "attributes": [
                    {"id": "attr_role_code", "name": "角色编码", "data_type": "string", "required": True, "constraints": {}},
                    {"id": "attr_role_name", "name": "角色名称", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_scope", "name": "范围", "data_type": "string", "required": False, "constraints": {}},
                ],
                "relations": [],
            },
            {
                "id": "obj_stage",
                "name": "阶段",
                "business_definition": "采购计划流程阶段",
                "stable_key": "attr_stage_id",
                "attributes": [
                    {"id": "attr_stage_id", "name": "阶段ID", "data_type": "string", "required": True, "constraints": {}},
                    {"id": "attr_stage_name", "name": "阶段名称", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_phase", "name": "阶段分组", "data_type": "string", "required": False, "constraints": {}},
                ],
                "relations": [],
            },
            {
                "id": "obj_supplier",
                "name": "供应商",
                "business_definition": "供应商主数据",
                "stable_key": "attr_supplier_code",
                "attributes": [
                    {"id": "attr_supplier_code", "name": "供应商编码", "data_type": "string", "required": True, "constraints": {}},
                    {"id": "attr_supplier_name", "name": "供应商名称", "data_type": "string", "required": False, "constraints": {}},
                    {"id": "attr_category", "name": "分类", "data_type": "string", "required": False, "constraints": {}},
                ],
                "relations": [],
            },
        ],
        "constraints": [],
    }


def _build_field_map_for_sheet(sheet_name: str) -> dict[str, str]:
    maps = {
        "Activity": {
            "activityCode": "attr_activity_code",
            "activityName": "attr_activity_name",
            "stageName": "attr_stage_name",
            "leadingRole": "attr_leading_role",
            "participantRole": "attr_participant_role",
            "inputDoc": "attr_input_doc",
            "outputDoc": "attr_output_doc",
            "description": "attr_description",
        },
        "Role": {
            "roleCode": "attr_role_code",
            "roleName": "attr_role_name",
            "scope": "attr_scope",
        },
        "Stage": {
            "stage_id": "attr_stage_id",
            "stageName": "attr_stage_name",
            "phase": "attr_phase",
        },
        "Supplier": {
            "supplierCode": "attr_supplier_code",
            "supplierName": "attr_supplier_name",
            "category": "attr_category",
        },
    }
    return maps.get(sheet_name, {})


def _verify_excel_rows(excel_path: Path) -> dict[str, int]:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    counts = {}
    for sn in wb.sheetnames:
        rows = list(wb[sn].iter_rows(values_only=True))
        counts[sn] = max(0, len(rows) - 1)
    return counts


def main() -> int:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        settings_module = __import__("config", fromlist=["get_settings"])
        settings = settings_module.get_settings()
        original_data_dir = settings.data_dir
        settings.data_dir = str(tmpdir_path / "data")

        original_call_llm = router_ontologies.call_llm
        router_ontologies.call_llm = _mock_call_llm

        client = TestClient(app)

        # 1. 创建本体
        r = client.post("/api/ontologies", params={"project_id": "l6_flow"}, json=_make_ontology_payload())
        assert r.status_code == 201, f"创建本体失败: {r.text}"
        ontology = r.json()["data"]
        ontology_id = ontology["id"]
        print(f"[OK] 创建本体 id={ontology_id}")

        # 2. 上传 L6 示例实例 Excel（显式字段映射）
        excel_path = _make_test_excel(tmpdir_path)
        expected_counts = _verify_excel_rows(excel_path)
        print(f"[INFO] Excel 各 sheet 数据行: {expected_counts}")

        with open(excel_path, "rb") as f:
            r = client.post(
                "/api/ingest/upload",
                params={"ontology_id": ontology_id, "project_id": "l6_flow"},
                data={"field_map": json.dumps(_build_field_map_for_sheet("Activity"))},
                files={"file": ("l6_instances.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert r.status_code == 201, f"上传失败: {r.text}"
        candidates = r.json()["data"]["candidates"]
        print(f"[OK] 上传生成候选 {len(candidates)} 组")

        # 3. 逐条审批
        total_imported = 0
        for cand in candidates:
            cid = cand["id"]
            r = client.post(f"/api/ingest/candidates/{cid}/approve", params={"project_id": "l6_flow"})
            assert r.status_code == 200, f"审批失败: {r.text}"
            result = r.json()["data"]
            print(f"[OK] candidate {cand['object_type']} imported={result['imported']} violations={len(result.get('violations', []))}")
            total_imported += result["imported"]

        print(f"[OK] 总计导入实例 {total_imported}")

        # 4. 验证 InstanceStore 落盘
        from core.instance_store import JsonlInstanceStore
        store = JsonlInstanceStore(project_id="l6_flow", data_dir=settings.data_dir)
        for sheet_name, expected in expected_counts.items():
            obj_id = f"obj_{sheet_name.lower()}"
            actual = store.count(object_type=obj_id)
            print(f"[CHECK] {obj_id}: 期望 {expected}, 实际 {actual}")
            assert actual == expected, f"{obj_id} 落盘数量不一致"

        # 5. 验证实例列表接口
        r = client.get("/api/ingest/instances", params={"project_id": "l6_flow", "ontology_id": ontology_id})
        assert r.status_code == 200, f"实例列表失败: {r.text}"
        items = r.json()["data"]["items"]
        print(f"[OK] 实例列表返回 {len(items)} 条")

        # 6. 测试 reject
        # 先造一个空候选：由于 preview 必填校验，直接用 upload 空 sheet 不可行，这里用 put 更新候选实例为空后 reject
        # 简化：验证 reject 接口可调用
        if candidates:
            r = client.post(f"/api/ingest/candidates/{candidates[0]['id']}/reject", params={"project_id": "l6_flow"})
            # 已 imported 的不能 reject，应返回 403
            assert r.status_code == 403, f"已导入候选应禁止拒绝: {r.text}"
            print("[OK] 已导入候选拒绝被禁止")

        # 7. 本体增强 preview（无需 LLM，构造 txt 文本）
        enrich_text = "本流程还包含审批记录对象，记录审批人、审批时间、审批意见。"
        r = client.post(
            f"/api/ontologies/{ontology_id}/enrich",
            params={"project_id": "l6_flow"},
            data={"instructions": "新增审批记录对象"},
            files={"file": ("supplement.txt", io.BytesIO(enrich_text.encode("utf-8")), "text/plain")},
        )
        assert r.status_code == 200, f"增强预览失败: {r.text}"
        summary = r.json()["data"]["summary"]
        print(f"[OK] 增强预览 summary={summary}")

        settings.data_dir = original_data_dir
        router_ontologies.call_llm = original_call_llm
        print("\n[PASS] L6 通用流程验证全部通过")
        return 0


if __name__ == "__main__":
    sys.exit(main())
